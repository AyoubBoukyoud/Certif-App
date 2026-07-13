"""Sheet / named-range / row-col / freeze / view checkers."""
from __future__ import annotations

import unicodedata

from openpyxl.utils import range_boundaries

from .base import CheckResult, as_bool, checker, fail, ok


def _norm(s):
    return unicodedata.normalize("NFC", str(s)).strip().casefold()


@checker("sheet_exists")
def check_sheet_exists(ctx, params) -> CheckResult:
    name = params["name"]
    if ctx.sheet_f(name) is not None:
        return ok(f"feuille «{name}» présente")
    return fail(f"feuille «{name}» absente (feuilles: {ctx.sheet_names()})")


@checker("sheet_visible")
def check_sheet_visible(ctx, params) -> CheckResult:
    name = params["name"]
    ws = ctx.sheet_f(name)
    if ws is None:
        return fail(f"feuille «{name}» introuvable")
    if ws.sheet_state == "visible":
        return ok(f"feuille «{name}» visible")
    return fail(f"feuille «{name}» état={ws.sheet_state}")


@checker("sheet_order")
def check_sheet_order(ctx, params) -> CheckResult:
    """Assert an index or relative ordering of sheets."""
    names = [_norm(n) for n in ctx.sheet_names()]
    name = _norm(params["name"])
    if name not in names:
        return fail(f"feuille «{params['name']}» absente")
    idx = names.index(name)

    if "index" in params:
        want = int(params["index"])
        want = want if want >= 0 else len(names) + want
        if idx != want:
            return fail(f"«{params['name']}» à l'index {idx} (attendu {want})")
    if params.get("before"):
        b = _norm(params["before"])
        if b in names and not idx < names.index(b):
            return fail(f"«{params['name']}» pas avant «{params['before']}»")
    if params.get("after"):
        a = _norm(params["after"])
        if a in names and not idx > names.index(a):
            return fail(f"«{params['name']}» pas après «{params['after']}»")
    return ok(f"«{params['name']}» bien positionnée (index {idx})")


@checker("sheet_view_flag")
def check_sheet_view_flag(ctx, params) -> CheckResult:
    """e.g. showFormulas / showGridLines on the first sheet view."""
    sheet = params["sheet"]
    ws = ctx.sheet_f(sheet)
    if ws is None:
        return fail(f"feuille «{sheet}» introuvable")
    flag = params.get("flag", "showFormulas")
    expected = as_bool(params.get("expected"), True)
    view = ws.sheet_view
    actual = bool(getattr(view, flag, False))
    if actual == expected:
        return ok(f"{sheet}: {flag}={actual}")
    return fail(f"{sheet}: {flag}={actual} ≠ {expected}")


@checker("named_range_empty")
def check_named_range_empty(ctx, params) -> CheckResult:
    name = params["name"]
    wb = ctx.wb_values
    dn = wb.defined_names.get(name)
    if dn is None:
        return fail(f"nom «{name}» introuvable")
    for title, coord in dn.destinations:
        ws = ctx._find_sheet(wb, title)
        if ws is None:
            continue
        for row in ws[coord]:
            for cell in row:
                if cell.value not in (None, ""):
                    return fail(f"«{name}» non vide en {title}!{cell.coordinate}")
    return ok(f"plage nommée «{name}» vidée")


@checker("hidden_rows")
def check_hidden_rows(ctx, params) -> CheckResult:
    sheet = params["sheet"]
    ws = ctx.sheet_f(sheet)
    if ws is None:
        return fail(f"feuille «{sheet}» introuvable")
    rows = [int(r) for r in str(params["rows"]).split(",")]
    for r in rows:
        if not ws.row_dimensions[r].hidden:
            return fail(f"{sheet}: ligne {r} non masquée")
    return ok(f"{sheet}: lignes {rows} masquées")


@checker("hidden_cols")
def check_hidden_cols(ctx, params) -> CheckResult:
    sheet = params["sheet"]
    ws = ctx.sheet_f(sheet)
    if ws is None:
        return fail(f"feuille «{sheet}» introuvable")
    cols = [c.strip() for c in str(params["cols"]).split(",")]
    for c in cols:
        if not ws.column_dimensions[c].hidden:
            return fail(f"{sheet}: colonne {c} non masquée")
    return ok(f"{sheet}: colonnes {cols} masquées")


@checker("freeze_panes")
def check_freeze_panes(ctx, params) -> CheckResult:
    sheet = params["sheet"]
    ws = ctx.sheet_f(sheet)
    if ws is None:
        return fail(f"feuille «{sheet}» introuvable")
    want = params["cell"]
    actual = ws.freeze_panes
    if actual and str(actual).replace("$", "") == want.replace("$", ""):
        return ok(f"{sheet}: volets figés en {actual}")
    return fail(f"{sheet}: volets figés = {actual} (attendu {want})")


@checker("column_widths_changed")
def check_column_widths_changed(ctx, params) -> CheckResult:
    """Auto-fit: widths of the given columns differ from the pristine file."""
    sheet = params["sheet"]
    ws = ctx.sheet_f(sheet)
    wp = ctx.sheet_p(sheet)
    if ws is None or wp is None:
        return fail(f"feuille «{sheet}» introuvable")
    min_c, _, max_c, _ = range_boundaries(params.get("range", "A1:A1"))
    from openpyxl.utils.cell import get_column_letter

    changed = 0
    for ci in range(min_c, max_c + 1):
        col = get_column_letter(ci)
        a = ws.column_dimensions[col].width
        b = wp.column_dimensions[col].width
        if (a or 0) != (b or 0):
            changed += 1
    if changed == 0:
        return fail(f"{sheet}: largeurs de colonnes inchangées")
    return ok(f"{sheet}: {changed} largeurs de colonnes ajustées")
