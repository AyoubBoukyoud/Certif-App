"""Cell / formula / value / hyperlink / merge checkers."""
from __future__ import annotations

import re

from openpyxl.utils import range_boundaries
from openpyxl.utils.cell import get_column_letter

from ..formula_locale import normalize
from .base import CheckResult, as_float, checker, fail, ok


def _iter_range_cells(ws, rng: str):
    min_c, min_r, max_c, max_r = range_boundaries(rng)
    for r in range(min_r, max_r + 1):
        for c in range(min_c, max_c + 1):
            yield ws.cell(row=r, column=c)


def _num_equal(a, b, tol: float) -> bool:
    try:
        return abs(float(a) - float(b)) <= tol
    except (TypeError, ValueError):
        return False


def _match_formula(formula: str, params: dict) -> str | None:
    """Return an error string if the formula fails the pattern/contains rule."""
    norm = normalize(formula)
    if not norm:
        return "cellule vide / sans formule"
    pattern = params.get("pattern")
    if pattern and not (re.search(pattern, norm, re.IGNORECASE)
                        or re.search(pattern, str(formula).upper(), re.IGNORECASE)):
        return f"formule «{formula}» ne correspond pas au motif attendu"
    contains = params.get("contains") or params.get("must_contain")
    if contains and normalize(contains) not in norm:
        return f"formule «{formula}» ne contient pas «{contains}»"
    return None


@checker("formula")
def check_formula(ctx, params) -> CheckResult:
    sheet, cell = params["sheet"], params["cell"]
    ws = ctx.sheet_f(sheet)
    if ws is None:
        return fail(f"feuille «{sheet}» introuvable")
    formula = ws[cell].value
    err = _match_formula(formula, params)
    if err:
        return fail(f"{sheet}!{cell}: {err}")

    expected = params.get("expected_value")
    if expected is not None:
        tol = as_float(params.get("tolerance"), 0.0)
        wsv = ctx.sheet_v(sheet)
        got = wsv[cell].value if wsv is not None else None
        if expected == "from_solution":
            sol = ctx.wb_solution
            if sol is None:
                return ok(f"{sheet}!{cell}: formule OK (valeur non vérifiée)")
            sws = ctx._find_sheet(sol, sheet)
            want = sws[cell].value if sws is not None else None
            if want is not None and not _num_equal(got, want, tol) and str(got) != str(want):
                return fail(f"{sheet}!{cell}: valeur {got!r} ≠ attendue {want!r}")
        else:
            if not _num_equal(got, expected, tol) and str(got) != str(expected):
                return fail(f"{sheet}!{cell}: valeur {got!r} ≠ {expected!r}")
    return ok(f"{sheet}!{cell}: formule conforme")


@checker("formula_column")
def check_formula_column(ctx, params) -> CheckResult:
    """Every non-empty data cell in a column/range must match the pattern."""
    sheet = params["sheet"]
    ws = ctx.sheet_f(sheet)
    if ws is None:
        return fail(f"feuille «{sheet}» introuvable")

    rng = params.get("range")
    if not rng:
        col = params.get("column")
        if params.get("header"):
            col = ctx.header_column(ws, params["header"],
                                    int(params.get("header_row", 1)))
        if not col:
            return fail("ni range ni colonne résolue pour formula_column")
        r1 = int(params.get("row_start", 2))
        r2 = int(params.get("row_end", ws.max_row))
        rng = f"{col}{r1}:{col}{r2}"

    checked = 0
    for c in _iter_range_cells(ws, rng):
        if c.value in (None, ""):
            continue
        err = _match_formula(c.value, params)
        if err:
            return fail(f"{sheet}!{c.coordinate}: {err}")
        checked += 1
    if checked == 0:
        return fail(f"{sheet}!{rng}: aucune formule trouvée")
    return ok(f"{sheet}!{rng}: {checked} formules conformes")


@checker("range_values")
def check_range_values(ctx, params) -> CheckResult:
    """Compare cached values of a range against the solution (or a literal list)."""
    sheet, rng = params["sheet"], params["range"]
    wsv = ctx.sheet_v(sheet)
    if wsv is None:
        return fail(f"feuille «{sheet}» introuvable")
    got = [c.value for c in _iter_range_cells(wsv, rng)]

    literal = params.get("equals")
    if literal is not None:
        want = [x.strip() for x in literal.split("|")]
        got_s = [("" if v is None else str(v)) for v in got]
        if got_s[: len(want)] != want:
            return fail(f"{sheet}!{rng}: valeurs {got_s} ≠ {want}")
        return ok(f"{sheet}!{rng}: valeurs conformes")

    sol = ctx.wb_solution
    if sol is None:
        non_empty = sum(1 for v in got if v not in (None, ""))
        if non_empty == 0:
            return fail(f"{sheet}!{rng}: plage vide")
        return ok(f"{sheet}!{rng}: {non_empty} valeurs présentes (solution absente)")
    sws = ctx._find_sheet(sol, sheet)
    want = [c.value for c in _iter_range_cells(sws, params.get("solution_range", rng))]
    if [str(x) for x in got] != [str(x) for x in want]:
        return fail(f"{sheet}!{rng}: valeurs ≠ solution")
    return ok(f"{sheet}!{rng}: valeurs = solution")


@checker("text_absent_present")
def check_text(ctx, params) -> CheckResult:
    """Assert some text is gone and/or some text now appears (find & replace)."""
    sheets = [params["sheet"]] if params.get("sheet") else ctx.sheet_names()
    absent = params.get("absent")
    present = params.get("present")
    min_count = int(params.get("min_count", 1))

    present_hits = 0
    for name in sheets:
        wsv = ctx.sheet_v(name)
        if wsv is None:
            continue
        for row in wsv.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                text = str(c.value)
                if absent and absent.lower() in text.lower():
                    return fail(f"«{absent}» encore présent en {name}!{c.coordinate}")
                if present and present.lower() in text.lower():
                    present_hits += 1
    if present and present_hits < min_count:
        return fail(f"«{present}» trouvé {present_hits}× (attendu ≥ {min_count})")
    return ok("texte conforme")


@checker("hyperlink")
def check_hyperlink(ctx, params) -> CheckResult:
    sheet, cell = params["sheet"], params["cell"]
    ws = ctx.sheet_f(sheet)
    if ws is None:
        return fail(f"feuille «{sheet}» introuvable")
    link = ws[cell].hyperlink
    if link is None:
        return fail(f"{sheet}!{cell}: aucun lien hypertexte")
    target = params.get("target")
    location = params.get("location")
    if target:
        actual = link.target or ""
        if not re.search(target, actual, re.IGNORECASE):
            return fail(f"{sheet}!{cell}: cible «{actual}» ≠ «{target}»")
    if location:
        actual = (link.location or "").replace("'", "")
        if actual.casefold() != location.replace("'", "").casefold():
            return fail(f"{sheet}!{cell}: emplacement «{actual}» ≠ «{location}»")
    return ok(f"{sheet}!{cell}: lien conforme")


@checker("merged_range")
def check_merged(ctx, params) -> CheckResult:
    sheet, rng = params["sheet"], params["range"]
    ws = ctx.sheet_f(sheet)
    if ws is None:
        return fail(f"feuille «{sheet}» introuvable")
    merged = {str(m) for m in ws.merged_cells.ranges}
    if rng not in merged:
        return fail(f"{sheet}: plage «{rng}» non fusionnée (fusions: {sorted(merged)})")
    return ok(f"{sheet}: «{rng}» fusionnée")


@checker("number_format")
def check_number_format(ctx, params) -> CheckResult:
    sheet, rng = params["sheet"], params["range"]
    ws = ctx.sheet_f(sheet)
    if ws is None:
        return fail(f"feuille «{sheet}» introuvable")
    want = params.get("contains")           # e.g. "0.000"
    exact = params.get("format")
    checked = 0
    for c in _iter_range_cells(ws, rng):
        if c.value in (None, ""):
            continue
        fmt = c.number_format or ""
        if exact and fmt != exact:
            return fail(f"{sheet}!{c.coordinate}: format «{fmt}» ≠ «{exact}»")
        if want and want not in fmt:
            return fail(f"{sheet}!{c.coordinate}: format «{fmt}» ne contient pas «{want}»")
        checked += 1
    if checked == 0:
        return fail(f"{sheet}!{rng}: aucune cellule à vérifier")
    return ok(f"{sheet}!{rng}: format conforme ({checked} cellules)")
