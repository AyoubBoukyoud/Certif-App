"""Shared state handed to every checker for one project submission.

Loads the candidate workbook twice (formulas + cached values), the pristine
workbook (for "unchanged"/diff checks), an optional solution workbook (for
value-based expectations) and keeps the raw ZIP handy for XML-level checks
that openpyxl doesn't surface (charts, sparklines, drawings, doc props).
"""
from __future__ import annotations

import unicodedata
import warnings
import zipfile
from functools import cached_property
from pathlib import Path

from openpyxl import load_workbook

from .. import config

warnings.filterwarnings("ignore", module="openpyxl")


def _nfc(s: str) -> str:
    return unicodedata.normalize("NFC", s)


class GradingContext:
    def __init__(self, user_path: str | Path, pristine_path: str | Path,
                 solution_path: str | Path | None = None) -> None:
        self.user_path = Path(user_path)
        self.pristine_path = Path(pristine_path)
        self.solution_path = Path(solution_path) if solution_path else None

    # -- workbooks (lazy) --------------------------------------------------

    @cached_property
    def wb_formulas(self):
        return load_workbook(self.user_path, data_only=False)

    @cached_property
    def wb_values(self):
        return load_workbook(self.user_path, data_only=True)

    @cached_property
    def wb_pristine(self):
        return load_workbook(self.pristine_path, data_only=False)

    @cached_property
    def wb_solution(self):
        if self.solution_path and self.solution_path.exists():
            return load_workbook(self.solution_path, data_only=True)
        return None

    # -- sheet resolution (accent/space tolerant) --------------------------

    @staticmethod
    def _find_sheet(wb, name: str):
        target = _nfc(name).strip().casefold()
        for ws in wb.worksheets:
            if _nfc(ws.title).strip().casefold() == target:
                return ws
        return None

    def sheet_f(self, name: str):
        return self._find_sheet(self.wb_formulas, name)

    def sheet_v(self, name: str):
        return self._find_sheet(self.wb_values, name)

    def sheet_p(self, name: str):
        return self._find_sheet(self.wb_pristine, name)

    def sheet_names(self):
        return [ws.title for ws in self.wb_formulas.worksheets]

    # -- raw XML -----------------------------------------------------------

    @cached_property
    def zip(self) -> zipfile.ZipFile:
        return zipfile.ZipFile(self.user_path)

    def read_xml(self, arc_name: str) -> bytes | None:
        try:
            return self.zip.read(arc_name)
        except KeyError:
            return None

    def namelist(self) -> list[str]:
        return self.zip.namelist()

    # -- helpers -----------------------------------------------------------

    @staticmethod
    def get_table(ws, table_name: str):
        for name in getattr(ws, "tables", {}):
            if name.casefold() == table_name.casefold():
                return ws.tables[name]
        return None

    @staticmethod
    def header_column(ws, header: str, header_row: int = 1) -> str | None:
        """Return the column letter whose header cell matches ``header``."""
        want = _nfc(str(header)).strip().casefold()
        for cell in ws[header_row]:
            if cell.value is not None and _nfc(str(cell.value)).strip().casefold() == want:
                return cell.column_letter
        return None

    def close(self) -> None:
        try:
            self.zip.close()
        except Exception:
            pass
