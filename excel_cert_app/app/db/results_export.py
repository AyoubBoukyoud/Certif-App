"""Append one row per finished session to a shared results workbook.

Robust against the file being open/locked (falls back to a timestamped copy)
so a candidate is never blocked by an admin having the results file open.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

HEADER = [
    "Date", "Prénom", "Nom", "CIN", "Date de naissance", "Établissement",
    "Variante", "Durée (min)", "Score %", "Résultat",
]


def _load_or_create(path: Path) -> Workbook:
    if path.exists():
        return load_workbook(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Résultats"
    ws.append(HEADER)
    for cell in ws[1]:
        cell.font = cell.font.copy(bold=True)
    return wb


def append_result(
    results_path: Path | str,
    *,
    first_name: str,
    last_name: str,
    cin: str,
    birth_date: str,
    institution: str,
    variant: str,
    duration_s: int,
    score_pct: float,
    passed: bool,
) -> Path:
    """Append a result row. Returns the path actually written (may be a fallback)."""
    results_path = Path(results_path)
    results_path.parent.mkdir(parents=True, exist_ok=True)

    row = [
        datetime.now().strftime("%Y-%m-%d %H:%M"),
        first_name, last_name, cin, birth_date, institution, variant,
        round(duration_s / 60, 1),
        round(score_pct, 1),
        "RÉUSSI" if passed else "ÉCHOUÉ",
    ]

    def _write(target: Path) -> None:
        wb = _load_or_create(target)
        wb["Résultats"].append(row)
        wb.save(target)

    try:
        _write(results_path)
        return results_path
    except (PermissionError, OSError):
        # File is likely open in Excel; write a timestamped fallback instead.
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        fallback = results_path.with_name(f"{results_path.stem}_{stamp}.xlsx")
        _write(fallback)
        return fallback
