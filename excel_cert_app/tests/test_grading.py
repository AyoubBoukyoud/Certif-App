"""Calibration-style grading tests.

For each checker we build a 'candidate' workbook from the pristine file with a
single edit applied and assert the corresponding task flips from FAIL (on the
pristine file) to PASS. This is the core quality gate for the grading engine.
"""
import os
import tempfile
import warnings

os.environ.setdefault("EXCEL_CERT_DATA_DIR", tempfile.mkdtemp(prefix="excel_cert_test_"))
warnings.filterwarnings("ignore")

from openpyxl import load_workbook  # noqa: E402

from app import config, content  # noqa: E402
from app.grading.engine import GradingEngine  # noqa: E402

ENGINE = GradingEngine()


def _grade_task(project_no: int, task_id: int, workbook_path):
    project = content.load_project(project_no)
    outcomes = ENGINE.grade_project(project, workbook_path)
    return outcomes[task_id]


def _modify(project_no: int, edit_fn):
    """Copy the pristine workbook, apply edit_fn(wb), save to a temp file."""
    src = config.project_workbook(project_no)
    wb = load_workbook(src)
    edit_fn(wb)
    tmp = tempfile.mktemp(suffix=".xlsx")
    wb.save(tmp)
    return tmp


def test_pristine_files_fail_their_tasks():
    # A handful of representative tasks must fail on the un-worked file.
    for pno, tid in [(1, 1), (1, 3), (3, 2), (8, 1), (12, 1), (6, 1)]:
        ok, _ = _grade_task(pno, tid, config.project_workbook(pno))
        assert not ok, f"P{pno} T{tid} should fail on pristine"


def test_banded_rows_p1_t1():
    def edit(wb):
        t = wb["Étudiants"].tables["Tableau1"]
        t.tableStyleInfo.showRowStripes = True
    ok, det = _grade_task(1, 1, _modify(1, edit))
    assert ok, det


def test_formula_average_p1_t3():
    def edit(wb):
        wb["Étudiants"]["N2"] = "=AVERAGE(K2:K501)"
    ok, det = _grade_task(1, 3, _modify(1, edit))
    assert ok, det


def test_hidden_rows_p3_t2():
    def edit(wb):
        ws = wb["Fournitures scolaires"]
        ws.row_dimensions[29].hidden = True
        ws.row_dimensions[30].hidden = True
    ok, det = _grade_task(3, 2, _modify(3, edit))
    assert ok, det


def test_freeze_panes_p3_t3():
    def edit(wb):
        wb["Fournitures scolaires"].freeze_panes = "A10"
    ok, det = _grade_task(3, 3, _modify(3, edit))
    assert ok, det


def test_named_range_cleared_p2_t2():
    def edit(wb):
        dn = wb.defined_names.get("Génie")
        for title, coord in dn.destinations:
            ws = wb[title]
            for row in ws[coord]:
                for cell in row:
                    cell.value = None
    ok, det = _grade_task(2, 2, _modify(2, edit))
    assert ok, det


def test_new_sheet_p2_t3():
    def edit(wb):
        wb.create_sheet("Contributeurs -10")
    ok, det = _grade_task(2, 3, _modify(2, edit))
    assert ok, det


def test_hidden_cols_p12_t1():
    def edit(wb):
        wb["Academy"].column_dimensions["D"].hidden = True
    ok, det = _grade_task(12, 1, _modify(12, edit))
    assert ok, det


def test_merged_range_p8_t1():
    def edit(wb):
        wb["Ingrédients"].merge_cells("A1:N1")
    ok, det = _grade_task(8, 1, _modify(8, edit))
    assert ok, det


def test_countif_formula_p6_t1():
    def edit(wb):
        wb["Les Forums"]["K10"] = '=COUNTIF(G8:G25,">=12")'
    ok, det = _grade_task(6, 1, _modify(6, edit))
    assert ok, det


def test_lowercase_formula_column_p12_t2():
    def edit(wb):
        ws = wb["Academy"]
        for r in range(8, 70):
            ws[f"L{r}"] = f"=LOWER(A{r})"
    ok, det = _grade_task(12, 2, _modify(12, edit))
    assert ok, det


def test_print_titles_p12_t4():
    def edit(wb):
        wb["Academy"].print_title_rows = "$7:$7"
    ok, det = _grade_task(12, 4, _modify(12, edit))
    assert ok, det


def test_find_replace_p4_t2():
    def edit(wb):
        ws = wb["Achat de moteurs"]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and "Pick-up" in c.value:
                    c.value = c.value.replace("Pick-up", "Caravane")
    ok, det = _grade_task(4, 2, _modify(4, edit))
    assert ok, det
